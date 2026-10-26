#!/usr/bin/env python3
"""
scripts/udemy_progress_sync.py
==============================
Udemy Course Progress Synchronization Tool for Excel VBA Second Brain.

Synchronizes course progress into the Obsidian Second Brain via three methods:
  1. From Excel Outline:
     python scripts/udemy_progress_sync.py --from-excel [path/to/Detailed_VBA_Course_Outline.xlsx]
  2. From Udemy Browser Export (JSON):
     python scripts/udemy_progress_sync.py --from-json udemy_progress.json
  3. Direct CLI Lecture Flag:
     python scripts/udemy_progress_sync.py --complete S05-01
     python scripts/udemy_progress_sync.py --section 5

After updating config/course_manifest.json, this script automatically runs
scripts/sync_vault.py to refresh all dashboards and roadmaps across the vault.
"""

import os
import sys
import json
import argparse
import re

sys.stdout.reconfigure(encoding='utf-8')

ROOT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MANIFEST_PATH = os.path.join(ROOT_DIR, "config", "course_manifest.json")
DEFAULT_EXCEL_PATH = r"D:\courses\Data Science\Data Engineering\Udemy Courses\VBA & Marcos Course\Detailed_VBA_Course_Outline.xlsx"

def load_manifest():
    with open(MANIFEST_PATH, 'r', encoding='utf-8') as f:
        return json.load(f)

def save_manifest(manifest):
    with open(MANIFEST_PATH, 'w', encoding='utf-8') as f:
        json.dump(manifest, f, indent=2, ensure_ascii=False)

def update_lecture_frontmatter(lecture_id, is_completed):
    lectures_dir = os.path.join(ROOT_DIR, "01 - Course", "Lectures")
    for fn in os.listdir(lectures_dir):
        if fn.startswith(f"Lecture {lecture_id}"):
            fpath = os.path.join(lectures_dir, fn)
            with open(fpath, 'r', encoding='utf-8') as fp:
                content = fp.read()
            new_status = "completed" if is_completed else "not-started"
            updated = re.sub(r'status:\s*(?:completed|not-started|planned)', f'status: {new_status}', content)
            if updated != content:
                with open(fpath, 'w', encoding='utf-8') as fp:
                    fp.write(updated)
            break

def sync_from_excel(excel_path):
    print(f"📖 Reading progress from Excel outline: {excel_path}")
    import openpyxl
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb['VBA Course Overview']
    
    manifest = load_manifest()
    
    current_sec = None
    lec_counter = 0
    updated_count = 0
    
    for r in range(3, ws.max_row + 1):
        sec_val = ws.cell(r, 1).value
        title = ws.cell(r, 3).value
        completed_val = ws.cell(r, 5).value
        notes_val = ws.cell(r, 6).value
        
        if sec_val is not None and str(sec_val).strip().isdigit():
            current_sec = int(sec_val)
            lec_counter = 0
        elif current_sec is not None and title is not None and str(title).strip():
            lec_counter += 1
            lec_id = f"S{current_sec:02d}-{lec_counter:02d}"
            
            is_done = False
            if current_sec in [1, 2, 3, 4] or notes_val == "Done" or completed_val in [True, "True", "true", 1]:
                is_done = True
                
            # Find in manifest
            for sec in manifest["sections"]:
                if sec["section_number"] == current_sec:
                    for lec in sec["lectures"]:
                        if lec["lecture_id"] == lec_id:
                            if lec["completed"] != is_done or lec["user_notes"] != str(notes_val or "").strip():
                                lec["completed"] = is_done
                                if notes_val:
                                    lec["user_notes"] = str(notes_val).strip()
                                updated_count += 1
                            update_lecture_frontmatter(lec_id, is_done)
                            
    save_manifest(manifest)
    print(f"✓ Synchronized {updated_count} lecture records from Excel.")

def sync_from_json(json_path):
    print(f"📥 Ingesting progress from JSON: {json_path}")
    with open(json_path, 'r', encoding='utf-8') as fp:
        data = json.load(fp)
        
    # Expecting list of completed lecture titles or IDs: ["S05-01", ...] or [{"id": "S05-01", "completed": true}]
    manifest = load_manifest()
    completed_ids = set()
    
    if isinstance(data, list):
        for item in data:
            if isinstance(item, str):
                completed_ids.add(item.strip())
            elif isinstance(item, dict) and item.get("completed"):
                completed_ids.add(item.get("id") or item.get("lecture_id") or item.get("title"))
                
    updated_count = 0
    for sec in manifest["sections"]:
        for lec in sec["lectures"]:
            if lec["lecture_id"] in completed_ids or lec["title"] in completed_ids:
                if not lec["completed"]:
                    lec["completed"] = True
                    updated_count += 1
                    update_lecture_frontmatter(lec["lecture_id"], True)
                    
    save_manifest(manifest)
    print(f"✓ Marked {updated_count} newly completed lectures from JSON export.")

def mark_single_lecture(lecture_id):
    manifest = load_manifest()
    found = False
    for sec in manifest["sections"]:
        for lec in sec["lectures"]:
            if lec["lecture_id"].lower() == lecture_id.lower():
                lec["completed"] = True
                update_lecture_frontmatter(lec["lecture_id"], True)
                found = True
                break
    if found:
        save_manifest(manifest)
        print(f"✓ Marked lecture {lecture_id} as completed.")
    else:
        print(f"❌ Lecture {lecture_id} not found in manifest.")

def mark_section(section_num):
    manifest = load_manifest()
    found = False
    for sec in manifest["sections"]:
        if sec["section_number"] == section_num:
            for lec in sec["lectures"]:
                lec["completed"] = True
                update_lecture_frontmatter(lec["lecture_id"], True)
            sec["status"] = "completed"
            found = True
            break
    if found:
        save_manifest(manifest)
        print(f"✓ Marked entire Section {section_num:02d} as completed.")
    else:
        print(f"❌ Section {section_num} not found in manifest.")

def run_vault_sync():
    # Import sync_vault module directly
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    import sync_vault
    sync_vault.main()

def main():
    parser = argparse.ArgumentParser(description="Synchronize Udemy course progress with Obsidian Second Brain.")
    parser.add_argument("--from-excel", nargs="?", const=DEFAULT_EXCEL_PATH, help="Sync from Detailed_VBA_Course_Outline.xlsx")
    parser.add_argument("--from-json", help="Sync from exported udemy_progress.json")
    parser.add_argument("--complete", help="Mark single lecture complete by ID (e.g., S05-01)")
    parser.add_argument("--section", type=int, help="Mark entire section complete by number (e.g., 5)")
    
    args = parser.parse_args()
    
    if args.from_excel:
        sync_from_excel(args.from_excel)
    elif args.from_json:
        sync_from_json(args.from_json)
    elif args.complete:
        mark_single_lecture(args.complete)
    elif args.section:
        mark_section(args.section)
    else:
        # Default: sync from local Excel outline
        sync_from_excel(DEFAULT_EXCEL_PATH)
        
    run_vault_sync()

if __name__ == "__main__":
    main()
